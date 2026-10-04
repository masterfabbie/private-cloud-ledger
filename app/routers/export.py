import csv
import io
import json
from calendar import month_name
from datetime import date

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from fastapi.responses import Response
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import models
from app.auth import get_current_user
from app.config import get_settings
from app.db import get_db
from app.services import backup
from app.services.queries import TxFilters, filtered_transactions, tx_filters

router = APIRouter(prefix="/api/export", tags=["export"])

HEADERS = ["Date", "Description", "Payer/Payee", "Amount", "Type", "Category", "Tags", "Account", "IBAN", "Notes"]


def _filename(f: TxFilters, ext: str) -> str:
    # Same naming as the original tracker: transactions_<year>_<Month>_<today>.csv
    name = "transactions"
    if f.year:
        name += f"_{f.year}"
    if f.month:
        name += f"_{month_name[f.month]}"
    return f"{name}_{date.today().isoformat()}.{ext}"


def _rows(db: Session, user: models.User, f: TxFilters):
    """One row per transaction, or one row per part for split transactions (the amounts add up)."""
    for t in db.scalars(filtered_transactions(user.id, f)).unique():
        lines = [(t.amount_cents, t.category, "")] if not t.splits else [
            (sp.amount_cents, sp.category, sp.note) for sp in t.splits
        ]
        for cents, category, note in lines:
            yield t, cents, [
                t.booking_date.isoformat(),
                f"{t.description} ({note})" if note else t.description,
                t.payer,
                abs(cents) / 100,
                "income" if cents >= 0 else "expense",
                category.name if category else "",
                "; ".join(sorted(tag.name for tag in t.tags)),
                t.account.name,
                t.counterparty_iban,
                t.notes,
            ]


def _safe_cell(value):
    # Avoid spreadsheet formula injection from bank-provided text.
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


@router.get("/csv")
def export_csv(f: TxFilters = Depends(tx_filters), db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    writer.writerow(HEADERS)
    for _, _, row in _rows(db, user, f):
        row[3] = f"{row[3]:.2f}".replace(".", ",")  # European decimal comma, like the original export
        writer.writerow([_safe_cell(c) for c in row])
    return Response(
        "﻿" + buf.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{_filename(f, "csv")}"'},
    )


@router.get("/xlsx")
def export_xlsx(f: TxFilters = Depends(tx_filters), db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(HEADERS)
    for t, cents, row in _rows(db, user, f):
        row[0] = t.booking_date
        row[3] = cents / 100  # signed in Excel so SUM() works
        ws.append([_safe_cell(c) for c in row])
    for cell in ws["A"][1:]:
        cell.number_format = "DD.MM.YYYY"
    for cell in ws["D"][1:]:
        cell.number_format = "#,##0.00 €"
    ws.freeze_panes = "A2"
    for col, width in zip("ABCDEFGHIJ", (12, 50, 30, 12, 9, 18, 20, 16, 26, 30)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{_filename(f, "xlsx")}"'},
    )


@router.get("/json")
def export_json(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    """Full backup of the current user's data (restore it on the Settings page)."""
    body = json.dumps(backup.export_data(db, user), default=str, ensure_ascii=False, indent=1)
    return Response(
        body,
        media_type="application/json",
        headers={"Content-Disposition": f'attachment; filename="proud-ledger_backup_{date.today().isoformat()}.json"'},
    )


@router.post("/restore", tags=["backup"])
async def restore_backup(
    file: UploadFile = File(...),
    dry_run: bool = False,
    db: Session = Depends(get_db),
    user: models.User = Depends(get_current_user),
):
    """Check a backup file (dry_run=true) or replace all of the user's data with it."""
    limit = max(50, get_settings().max_upload_mb) * 1024 * 1024
    raw = await file.read(limit + 1)
    if len(raw) > limit:
        raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, "The backup file is too large")
    try:
        data = json.loads(raw.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "This is not a Proud Ledger backup file (not valid JSON).") from exc
    try:
        bk = backup.parse_backup(data)
    except backup.BackupError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(exc)) from exc
    result = {"backup": backup.summarize(bk), "current": backup.current_counts(db, user.id), "restored": False}
    if not dry_run:
        backup.restore(db, user, bk)
        result["restored"] = True
    return result
